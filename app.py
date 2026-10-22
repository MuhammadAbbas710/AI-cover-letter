"""AI Cover Letter Generator
Upload a resume, paste a job description, and get a personalized cover letter.
Built with Streamlit and Google Gemini Flash.
"""

import io
import os

import streamlit as st
from google import genai
from google.genai import types
from google.genai import errors as genai_errors
from pypdf import PdfReader
from docx import Document

# Primary Flash model, with a fallback if the primary is unavailable (HTTP 404).
# You can override the primary with a GEMINI_MODEL secret or environment variable.
DEFAULT_MODEL = "gemini-3.5-flash"
FALLBACK_MODEL = "gemini-2.5-flash"
MAX_RESUME_CHARS = 15000
MAX_JOB_CHARS = 8000

st.set_page_config(
    page_title="AI Cover Letter Generator",
    page_icon="✉️",
    layout="centered",
)


# ----------------------------- Helpers ------------------------------------ #
def get_api_key(user_key: str = "") -> str:
    """Return the API key from the sidebar, Streamlit secrets, or environment."""
    if user_key and user_key.strip():
        return user_key.strip()
    try:
        secret = st.secrets.get("GEMINI_API_KEY", "")
        if secret:
            return secret
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def extract_text(uploaded_file) -> str:
    """Extract plain text from an uploaded PDF, DOCX, or TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")


def build_prompt(resume, job, tone, length, company, role, name, extra) -> str:
    length_map = {
        "Short (~200 words)": "about 200 words",
        "Standard (~300 words)": "about 300 words",
        "Detailed (~400 words)": "about 400 words",
    }
    lines = [
        "You are an expert career coach and professional cover letter writer.",
        "Write a personalized cover letter using ONLY the facts in the resume below.",
        "",
        "Rules:",
        "- Do not invent employers, job titles, degrees, dates, metrics, or skills "
        "that are not in the resume.",
        "- Match the candidate's most relevant experience and skills to the "
        "requirements in the job description, using concrete examples from the resume.",
        "- Mirror important keywords from the job description naturally.",
        f"- Tone: {tone}.",
        f"- Length: {length_map[length]}.",
        "- Use a standard business letter structure: greeting, opening paragraph, "
        "1-2 body paragraphs, closing paragraph with a call to action, sign-off.",
        "- If the hiring manager's name is unknown, use 'Dear Hiring Manager,'.",
        "- Output only the final letter, with no commentary, notes, or markdown "
        "formatting. Do not leave bracketed placeholders unless information is "
        "truly missing from the resume.",
        "",
    ]
    if name:
        lines.append(f"Candidate name: {name}")
    if company:
        lines.append(f"Company: {company}")
    if role:
        lines.append(f"Role: {role}")
    if extra:
        lines.append(f"Additional instructions from the candidate: {extra}")
    lines += [
        "",
        "=== RESUME ===",
        resume[:MAX_RESUME_CHARS],
        "",
        "=== JOB DESCRIPTION ===",
        job[:MAX_JOB_CHARS],
    ]
    return "\n".join(lines)


def get_model_name() -> str:
    try:
        secret = st.secrets.get("GEMINI_MODEL", "")
        if secret:
            return secret
    except Exception:
        pass
    return os.environ.get("GEMINI_MODEL", DEFAULT_MODEL)


def generate_letter(api_key: str, prompt: str) -> str:
    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(temperature=0.7, max_output_tokens=2048)

    candidates = [get_model_name()]
    if FALLBACK_MODEL not in candidates:
        candidates.append(FALLBACK_MODEL)

    response = None
    last_error = None
    for model_name in candidates:
        try:
            response = client.models.generate_content(
                model=model_name, contents=prompt, config=config
            )
            break
        except genai_errors.APIError as e:
            last_error = e
            if getattr(e, "code", None) == 404:
                continue  # model not available, try the fallback
            raise
    if response is None:
        raise last_error

    text = getattr(response, "text", None)
    if not text or not text.strip():
        raise RuntimeError(
            "The model returned no text (the response may have been blocked). "
            "Try adjusting your inputs and generating again."
        )
    return text.strip()


def make_docx(letter: str) -> bytes:
    doc = Document()
    for block in letter.split("\n\n"):
        block = block.strip()
        if block:
            doc.add_paragraph(block)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------- UI --------------------------------------- #
st.title("✉️ AI Cover Letter Generator")
st.caption("Upload your resume, paste a job description, and get a tailored cover letter.")

with st.sidebar:
    st.header("Settings")
    user_key = st.text_input(
        "Gemini API key",
        type="password",
        help="Get a free key at https://aistudio.google.com/app/apikey. "
        "Leave blank if the app owner has configured one.",
    )
    tone = st.selectbox(
        "Tone",
        ["Professional", "Enthusiastic", "Confident", "Friendly", "Formal"],
    )
    length = st.selectbox(
        "Length",
        ["Short (~200 words)", "Standard (~300 words)", "Detailed (~400 words)"],
        index=1,
    )
    st.markdown("---")
    st.caption("Your resume is sent to Google's Gemini API to generate the letter. "
               "This app does not store it.")

resume_file = st.file_uploader("1. Upload your resume", type=["pdf", "docx", "txt"])
job_description = st.text_area(
    "2. Paste the job description", height=250,
    placeholder="Paste the full job posting here...",
)

col1, col2 = st.columns(2)
with col1:
    company = st.text_input("Company name (optional)")
with col2:
    role = st.text_input("Job title (optional)")
name = st.text_input("Your name (optional)", help="Used in the letter's sign-off.")
extra = st.text_area(
    "Anything else to include? (optional)", height=80,
    placeholder="e.g. mention I'm relocating, or emphasize my leadership experience",
)

if st.button("Generate cover letter", type="primary"):
    api_key = get_api_key(user_key)
    if not api_key:
        st.error("Please enter your Gemini API key in the sidebar.")
    elif resume_file is None:
        st.error("Please upload your resume.")
    elif not job_description.strip():
        st.error("Please paste the job description.")
    else:
        try:
            with st.spinner("Reading your resume..."):
                resume_text = extract_text(resume_file)
            if len(resume_text) < 50:
                st.error(
                    "Couldn't read enough text from the resume. If it's a scanned "
                    "image PDF, upload a text-based PDF or a DOCX instead."
                )
            else:
                prompt = build_prompt(
                    resume_text, job_description.strip(), tone, length,
                    company.strip(), role.strip(), name.strip(), extra.strip(),
                )
                with st.spinner("Writing your cover letter..."):
                    letter = generate_letter(api_key, prompt)
                    st.session_state["letter"] = letter
                    st.session_state["letter_editor"] = letter
        except ValueError as e:
            st.error(str(e))
        except Exception as e:
            st.error(f"Something went wrong: {e}")

if st.session_state.get("letter"):
    st.subheader("Your cover letter")
    edited = st.text_area("Edit as needed", height=450, key="letter_editor")
    d1, d2 = st.columns(2)
    with d1:
        st.download_button(
            "Download as .txt", data=edited, file_name="cover_letter.txt",
            mime="text/plain",
        )
    with d2:
        st.download_button(
            "Download as .docx", data=make_docx(edited),
            file_name="cover_letter.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
